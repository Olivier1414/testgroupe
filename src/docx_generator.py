"""
Générateur de documents Word (.docx) utilisant python-docx.
Produit des documents modifiables sous Word avec styles soignés,
numérotation automatique, corps personnalisé Jinja2 et tableaux récapitulatifs.
"""

from datetime import datetime
from pathlib import Path
from typing import Any, Dict

import docx
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from src.config import ORGANISATION_DEFAULT, OUTPUT_DIR


def appliquer_couleur_fond_cellule(cellule, couleur_hex: str) -> None:
    """Applique une couleur de fond (shading) à une cellule de tableau Word."""
    tc_pr = cellule._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), couleur_hex)
    tc_pr.append(shd)


def definir_bordures_cellule(cellule, couleur_hex: str = "CBD5E0", taille: str = "4") -> None:
    """Ajoute des bordures subtiles à une cellule."""
    tc_pr = cellule._tc.get_or_add_tcPr()
    tc_borders = OxmlElement("w:tcBorders")
    for bordure_nom in ("top", "left", "bottom", "right"):
        bordure = OxmlElement(f"w:{bordure_nom}")
        bordure.set(qn("w:val"), "single")
        bordure.set(qn("w:sz"), taille)
        bordure.set(qn("w:space"), "0")
        bordure.set(qn("w:color"), couleur_hex)
        tc_borders.append(bordure)
    tc_pr.append(tc_borders)


class GenerateurDOCX:
    """Génère des fichiers Word (.docx) conformes au cahier des charges."""

    def __init__(self, dossier_sortie: Path = OUTPUT_DIR):
        self.dossier_sortie = Path(dossier_sortie)
        self.dossier_sortie.mkdir(parents=True, exist_ok=True)

        # Palette de couleurs du projet
        self.c_primaire = RGBColor(0x1A, 0x36, 0x5D)    # #1A365D Bleu Marine
        self.c_secondaire = RGBColor(0x2B, 0x6C, 0xB0)  # #2B6CB0 Bleu Roi
        self.c_texte = RGBColor(0x2D, 0x37, 0x48)       # #2D3748 Gris Foncé
        self.c_muet = RGBColor(0x71, 0x80, 0x96)        # #718096 Gris Doux

    def generer(
        self,
        modele_rendu: Dict[str, str],
        contexte: Dict[str, Any],
        nom_fichier: str = None,
    ) -> Path:
        """
        Produit le document Word (.docx) sur disque.
        Retourne le chemin absolu du fichier généré.
        """
        numero = contexte.get("numero", "DOC-0001")
        if not nom_fichier:
            nom_propre = f"{contexte.get('nom', 'document')}_{contexte.get('prenom', '')}".strip("_")
            nom_fichier = f"{numero}_{nom_propre}.docx"

        chemin_docx = self.dossier_sortie / nom_fichier

        doc = Document()

        # Ajustement des marges (2 cm de chaque côté)
        for section in doc.sections:
            section.top_margin = Inches(0.8)
            section.bottom_margin = Inches(0.8)
            section.left_margin = Inches(0.8)
            section.right_margin = Inches(0.8)

            # Pied de page officiel
            footer = section.footer
            p_foot = footer.paragraphs[0]
            p_foot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            run_foot = p_foot.add_run(
                f"Document certifié {numero} - Généré automatiquement"
            )
            run_foot.font.name = "Calibri"
            run_foot.font.size = Pt(8.5)
            run_foot.font.color.rgb = self.c_muet

        # 1. En-tête : Tableau 2 colonnes (Organisme à gauche / Référence à droite)
        self._ajouter_en_tete(doc, contexte)

        # Ligne de séparation
        p_sep = doc.add_paragraph()
        p_sep.paragraph_format.space_before = Pt(4)
        p_sep.paragraph_format.space_after = Pt(14)
        run_sep = p_sep.add_run("―" * 54)
        run_sep.font.color.rgb = self.c_primaire
        run_sep.bold = True

        # 2. Titre et sous-titre
        titre = modele_rendu.get("titre", "DOCUMENT OFFICIEL")
        p_titre = doc.add_paragraph()
        p_titre.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_titre.paragraph_format.space_before = Pt(0)
        p_titre.paragraph_format.space_after = Pt(4)
        run_titre = p_titre.add_run(titre)
        run_titre.font.name = "Calibri"
        run_titre.font.size = Pt(20)
        run_titre.bold = True
        run_titre.font.color.rgb = self.c_primaire

        sous_titre = modele_rendu.get("sous_titre", "")
        if sous_titre:
            p_sub = doc.add_paragraph()
            p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_sub.paragraph_format.space_after = Pt(16)
            run_sub = p_sub.add_run(sous_titre)
            run_sub.font.name = "Calibri"
            run_sub.font.size = Pt(11)
            run_sub.italic = True
            run_sub.font.color.rgb = self.c_secondaire

        # 3. Encadré destinataire
        self._ajouter_cadre_destinataire(doc, contexte)

        # 4. Corps de texte (rendu par Jinja2)
        corps = modele_rendu.get("corps", "")
        for paragraphe in corps.split("\n\n"):
            para_nettoye = paragraphe.strip()
            if para_nettoye:
                p_corps = doc.add_paragraph()
                p_corps.paragraph_format.space_after = Pt(8)
                p_corps.paragraph_format.line_spacing = 1.15
                for idx, ligne in enumerate(para_nettoye.split("\n")):
                    if idx > 0:
                        p_corps.add_run("\n")
                    run_corps = p_corps.add_run(ligne)
                    run_corps.font.name = "Calibri"
                    run_corps.font.size = Pt(11)
                    run_corps.font.color.rgb = self.c_texte

        # 5. Tableau récapitulatif détaillé
        self._ajouter_tableau_details(doc, contexte)

        # 6. Bloc de signature
        signature_titre = modele_rendu.get("signature_titre", "La Direction")
        self._ajouter_bloc_signature(doc, contexte, signature_titre)

        # Sauvegarde
        doc.save(chemin_docx)
        return chemin_docx

    def _ajouter_en_tete(self, doc: Document, contexte: Dict[str, Any]) -> None:
        """Ajoute le tableau supérieur d'en-tête."""
        org = contexte.get("organisation", ORGANISATION_DEFAULT)
        numero = contexte.get("numero", "DOC-0001")
        date_emission = contexte.get("date", datetime.now().strftime("%d/%m/%Y"))

        table = doc.add_table(rows=1, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        cell_gauche = table.cell(0, 0)
        cell_gauche.width = Inches(4.2)
        p_org = cell_gauche.paragraphs[0]
        p_org.paragraph_format.space_after = Pt(0)
        r_org_nom = p_org.add_run(f"{org.get('nom', 'ORGANISME')}\n")
        r_org_nom.bold = True
        r_org_nom.font.size = Pt(10)
        r_org_nom.font.color.rgb = self.c_primaire

        r_org_details = p_org.add_run(
            f"{org.get('adresse', '')}\n"
            f"Tél : {org.get('telephone', '')} | Email : {org.get('email', '')}"
        )
        r_org_details.font.size = Pt(8.5)
        r_org_details.font.color.rgb = self.c_muet

        cell_droite = table.cell(0, 1)
        cell_droite.width = Inches(2.6)
        p_ref = cell_droite.paragraphs[0]
        p_ref.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_ref.paragraph_format.space_after = Pt(0)

        r_num = p_ref.add_run(f"RÉFÉRENCE : {numero}\n")
        r_num.bold = True
        r_num.font.size = Pt(10)
        r_num.font.color.rgb = self.c_primaire

        r_date = p_ref.add_run(f"Date d'émission : {date_emission}\nStatut : Certifié conforme")
        r_date.font.size = Pt(8.5)
        r_date.font.color.rgb = self.c_muet

    def _ajouter_cadre_destinataire(self, doc: Document, contexte: Dict[str, Any]) -> None:
        """Ajoute une grille récapitulative pour le bénéficiaire."""
        civilite = contexte.get("civilite", "")
        nom = contexte.get("nom", "").upper()
        prenom = contexte.get("prenom", "")
        email = contexte.get("email", "Non renseigné")
        date_ev = contexte.get("date", "")

        table = doc.add_table(rows=2, cols=4)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        donnees = [
            [("Destinataire :", True), (f"{civilite} {prenom} {nom}", False),
             ("Date :", True), (str(date_ev), False)],
            [("Contact :", True), (str(email), False),
             ("Réf. Dossier :", True), (str(contexte.get("numero", "")), False)],
        ]

        largeurs = [Inches(1.5), Inches(2.2), Inches(1.2), Inches(1.9)]

        for r_idx, row in enumerate(table.rows):
            for c_idx, cell in enumerate(row.cells):
                cell.width = largeurs[c_idx]
                appliquer_couleur_fond_cellule(cell, "F7FAFC")
                definir_bordures_cellule(cell, "E2E8F0", "4")

                libelle, est_gras = donnees[r_idx][c_idx]
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(2)
                run = p.add_run(libelle)
                run.font.name = "Calibri"
                run.font.size = Pt(9)
                run.bold = est_gras
                run.font.color.rgb = self.c_primaire if est_gras else self.c_texte

        p_spacer = doc.add_paragraph()
        p_spacer.paragraph_format.space_after = Pt(6)

    def _ajouter_tableau_details(self, doc: Document, contexte: Dict[str, Any]) -> None:
        """Ajoute un tableau récapitulatif pour les détails chiffrés ou qualitatifs."""
        lignes = []

        if contexte.get("motif"):
            lignes.append(("Objet / Motif :", str(contexte["motif"])))

        if contexte.get("montant") and str(contexte["montant"]) not in ("0", "0.00", ""):
            lignes.append(("Montant total :", f"{contexte['montant']} € (TTC)"))

        if contexte.get("note"):
            lignes.append(("Résultat d'évaluation :", f"{contexte['note']} / 20"))

        lignes.append(("Statut de conformité :", "Certifié et validé par l'organisme"))

        table = doc.add_table(rows=len(lignes) + 1, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        # En-tête du tableau
        hdr_cells = table.rows[0].cells
        hdr_cells[0].width = Inches(2.2)
        hdr_cells[1].width = Inches(4.6)
        for idx, titre in enumerate(["ÉLÉMENT", "DÉSIGNATION / VALEUR"]):
            appliquer_couleur_fond_cellule(hdr_cells[idx], "1A365D")
            p = hdr_cells[idx].paragraphs[0]
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run(titre)
            run.font.name = "Calibri"
            run.font.size = Pt(9.5)
            run.bold = True
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

        # Lignes de contenu
        for r_idx, (col1, col2) in enumerate(lignes, start=1):
            row_cells = table.rows[r_idx].cells
            row_cells[0].width = Inches(2.2)
            row_cells[1].width = Inches(4.6)

            fond = "F8FAFC" if r_idx % 2 == 0 else "FFFFFF"
            for c_idx, val in enumerate([col1, col2]):
                cell = row_cells[c_idx]
                appliquer_couleur_fond_cellule(cell, fond)
                definir_bordures_cellule(cell, "CBD5E0", "4")

                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                run = p.add_run(val)
                run.font.name = "Calibri"
                run.font.size = Pt(9)
                run.bold = (c_idx == 0)
                run.font.color.rgb = self.c_primaire if c_idx == 0 else self.c_texte

        p_spacer = doc.add_paragraph()
        p_spacer.paragraph_format.space_after = Pt(10)

    def _ajouter_bloc_signature(self, doc: Document, contexte: Dict[str, Any], signature_titre: str) -> None:
        """Ajoute un bloc de signature élégant cadré sur la droite."""
        date_emission = contexte.get("date", datetime.now().strftime("%d/%m/%Y"))

        table = doc.add_table(rows=1, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        cell_gauche = table.cell(0, 0)
        cell_gauche.width = Inches(3.6)

        cell_droite = table.cell(0, 1)
        cell_droite.width = Inches(3.2)
        appliquer_couleur_fond_cellule(cell_droite, "F8FAFC")
        definir_bordures_cellule(cell_droite, "CBD5E0", "4")

        p = cell_droite.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        r_fait = p.add_run(f"Fait le {date_emission}\n")
        r_fait.font.size = Pt(8.5)
        r_fait.italic = True
        r_fait.font.color.rgb = self.c_muet

        r_titre = p.add_run(f"{signature_titre}\n\n\n\n")
        r_titre.bold = True
        r_titre.font.size = Pt(10)
        r_titre.font.color.rgb = self.c_primaire

        r_cachet = p.add_run("[ Cachet et Signature autorisée ]")
        r_cachet.font.size = Pt(8.5)
        r_cachet.italic = True
        r_cachet.font.color.rgb = self.c_muet
