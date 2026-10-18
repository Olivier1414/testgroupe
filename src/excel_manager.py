"""
Module de gestion et de lecture des fichiers Excel (.xlsx) et CSV.
Permet d'importer des listes de personnes et de données pour générer des documents par lot.
Fournit également un utilitaire pour créer un fichier de test type.
"""

import csv
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional

try:
    import openpyxl
    OPENPYXL_DISPONIBLE = True
except ImportError:
    OPENPYXL_DISPONIBLE = False


COLONNES_ATTENDUES = [
    "type_document",
    "civilite",
    "nom",
    "prenom",
    "email",
    "date",
    "motif",
    "montant",
    "note",
]


def normaliser_cle(cle: str) -> str:
    """Nettoie et normalise le nom d'une colonne (minuscules, sans accents, sans espaces)."""
    if not cle:
        return ""
    cle = str(cle).strip().lower()
    remplacements = {
        "é": "e", "è": "e", "ê": "e", "ë": "e",
        "à": "a", "â": "a",
        "î": "i", "ï": "i",
        "ô": "o", "ö": "o",
        "ù": "u", "û": "u",
        "ç": "c",
        " ": "_", "-": "_", "'": "_",
    }
    for orig, dest in remplacements.items():
        cle = cle.replace(orig, dest)
    return cle


class GestionnaireExcel:
    """Lit les données depuis un fichier Excel ou CSV pour le publipostage."""

    @staticmethod
    def lire_fichier(chemin_fichier: Path) -> List[Dict[str, Any]]:
        """
        Détecte l'extension du fichier (.xlsx ou .csv) et extrait les lignes sous forme de dictionnaires.
        """
        chemin = Path(chemin_fichier)
        if not chemin.exists():
            raise FileNotFoundError(f"Le fichier spécifié n'existe pas : {chemin}")

        suffixe = chemin.suffix.lower()
        if suffixe == ".xlsx":
            return GestionnaireExcel._lire_xlsx(chemin)
        elif suffixe in (".csv", ".txt"):
            return GestionnaireExcel._lire_csv(chemin)
        else:
            raise ValueError(f"Format non supporté ({suffixe}). Utilisez un fichier .xlsx ou .csv.")

    @staticmethod
    def _lire_xlsx(chemin: Path) -> List[Dict[str, Any]]:
        """Lit un fichier Excel .xlsx via openpyxl."""
        if not OPENPYXL_DISPONIBLE:
            raise ImportError(
                "La bibliothèque 'openpyxl' n'est pas installée. "
                "Installez-la avec : pip install openpyxl"
            )

        wb = openpyxl.load_workbook(chemin, data_only=True)
        feuille = wb.active

        lignes = list(feuille.iter_rows(values_only=True))
        if not lignes:
            return []

        # En-têtes de colonnes
        en_tetes_bruts = lignes[0]
        en_tetes = [normaliser_cle(str(col)) if col is not None else f"col_{idx}" for idx, col in enumerate(en_tetes_bruts)]

        resultats = []
        for index_ligne, ligne in enumerate(lignes[1:], start=2):
            # Ignore les lignes entièrement vides
            if all(cellule is None or str(cellule).strip() == "" for cellule in ligne):
                continue

            ligne_dict: Dict[str, Any] = {}
            for col_idx, valeur in enumerate(ligne):
                nom_col = en_tetes[col_idx] if col_idx < len(en_tetes) else f"col_{col_idx}"
                if valeur is None:
                    valeur_str = ""
                elif isinstance(valeur, datetime):
                    valeur_str = valeur.strftime("%d/%m/%Y")
                else:
                    valeur_str = str(valeur).strip()

                ligne_dict[nom_col] = valeur_str

            # Garantir les clés de base
            ligne_dict = GestionnaireExcel._completer_cles_manquantes(ligne_dict)
            ligne_dict["_index_ligne"] = index_ligne
            resultats.append(ligne_dict)

        return resultats

    @staticmethod
    def _lire_csv(chemin: Path) -> List[Dict[str, Any]]:
        """Lit un fichier CSV avec détection de délimiteur (, ou ;)."""
        lignes = []
        # Essai de détection de l'encodage (utf-8 puis iso-8859-1)
        contenu = None
        for enc in ("utf-8-sig", "utf-8", "cp1252", "iso-8859-1"):
            try:
                with open(chemin, "r", encoding=enc) as f:
                    contenu = f.read()
                break
            except UnicodeDecodeError:
                continue

        if contenu is None:
            raise UnicodeError("Impossible de lire l'encodage du fichier CSV.")

        # Détection du délimiteur
        premiere_ligne = contenu.split("\n")[0] if contenu else ""
        delim = ";" if ";" in premiere_ligne else ","

        lecteur = csv.DictReader(contenu.splitlines(), delimiter=delim)
        for index_ligne, ligne in enumerate(lecteur, start=2):
            ligne_normalisee = {normaliser_cle(k): (v.strip() if v else "") for k, v in ligne.items()}
            ligne_normalisee = GestionnaireExcel._completer_cles_manquantes(ligne_normalisee)
            ligne_normalisee["_index_ligne"] = index_ligne
            lignes.append(ligne_normalisee)

        return lignes

    @staticmethod
    def _completer_cles_manquantes(donnees: Dict[str, Any]) -> Dict[str, Any]:
        """S'assure que toutes les colonnes standard sont présentes avec des valeurs par défaut."""
        valeurs_par_defaut = {
            "type_document": "attestation",
            "civilite": "M.",
            "nom": "Dupont",
            "prenom": "Jean",
            "email": "",
            "date": datetime.now().strftime("%d/%m/%Y"),
            "motif": "Session de formation",
            "montant": "0.00",
            "note": "15/20",
        }
        for k, v in valeurs_par_defaut.items():
            if k not in donnees or donnees[k] == "":
                donnees[k] = v
        return donnees

    @staticmethod
    def creer_donnees_exemples() -> List[Dict[str, Any]]:
        """Retourne une liste de données types pour démonstration."""
        return [
            {
                "type_document": "attestation",
                "civilite": "Mme",
                "nom": "Martin",
                "prenom": "Sophie",
                "email": "sophie.martin@email.fr",
                "date": "15/10/2026",
                "motif": "Formation Python pour débutants",
                "montant": "0.00",
                "note": "18",
            },
            {
                "type_document": "recu",
                "civilite": "M.",
                "nom": "Bernard",
                "prenom": "Thomas",
                "email": "t.bernard@example.com",
                "date": "16/10/2026",
                "motif": "Cotisation annuelle d'adhésion",
                "montant": "150.00",
                "note": "",
            },
            {
                "type_document": "convocation",
                "civilite": "Mme",
                "nom": "Dubois",
                "prenom": "Claire",
                "email": "claire.dubois@mail.com",
                "date": "25/11/2026",
                "motif": "Assemblée générale annuelle",
                "montant": "0.00",
                "note": "",
            },
            {
                "type_document": "bulletin",
                "civilite": "M.",
                "nom": "Lefebvre",
                "prenom": "Lucas",
                "email": "lucas.lefebvre@etudiant.fr",
                "date": "18/10/2026",
                "motif": "Algorithmique et automatisation de documents",
                "montant": "0.00",
                "note": "17.5",
            },
            {
                "type_document": "attestation",
                "civilite": "M.",
                "nom": "Moreau",
                "prenom": "Antoine",
                "email": "antoine.moreau@domaine.com",
                "date": "17/10/2026",
                "motif": "Stage pratique en développement logiciel",
                "montant": "0.00",
                "note": "16",
            },
        ]

    @staticmethod
    def generer_fichier_test(chemin_destination: Path) -> Path:
        """Génère un fichier de démonstration (.xlsx si openpyxl est disponible, sinon .csv)."""
        destination = Path(chemin_destination)
        destination.parent.mkdir(parents=True, exist_ok=True)
        donnees = GestionnaireExcel.creer_donnees_exemples()

        if destination.suffix.lower() == ".xlsx" and OPENPYXL_DISPONIBLE:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Documents"

            # En-têtes avec style
            ws.append(COLONNES_ATTENDUES)
            for col_num in range(1, len(COLONNES_ATTENDUES) + 1):
                cellule = ws.cell(row=1, column=col_num)
                cellule.font = openpyxl.styles.Font(bold=True, color="FFFFFF")
                cellule.fill = openpyxl.styles.PatternFill(
                    start_color="1A365D", end_color="1A365D", fill_type="solid"
                )
                cellule.alignment = openpyxl.styles.Alignment(horizontal="center")

            # Données
            for d in donnees:
                ws.append([d.get(col, "") for col in COLONNES_ATTENDUES])

            # Largeur automatique des colonnes
            for col in ws.columns:
                max_len = max(len(str(cell.value or "")) for cell in col)
                col_letter = openpyxl.utils.get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

            wb.save(destination)
            return destination
        else:
            # Fallback CSV
            dest_csv = destination.with_suffix(".csv")
            with open(dest_csv, "w", encoding="utf-8-sig", newline="") as f:
                writer = csv.DictWriter(f, fieldnames=COLONNES_ATTENDUES, delimiter=";")
                writer.writeheader()
                for d in donnees:
                    writer.writerow({col: d.get(col, "") for col in COLONNES_ATTENDUES})
            return dest_csv
